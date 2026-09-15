"""
ZH: 組織對照表的管理端編輯（v3.8）—— 學系→學院、行政單位、校區。

ZH: 端點（都掛在 /api/v1/admin/org 底下，全部要 admin）：
      GET  /departments   學系清單 + 每個系有幾個人
      PUT  /departments   批次寫入（新增／改學院／改校區／停用／改名）
      GET  /units         行政單位清單 + 每個單位有幾個人
      PUT  /units         同上

ZH: 🔴 **沒有刪除，只有停用。** `users.department` 存的是**系名本身**、
    `users.unit` 存的是 `org_units.path` —— 都不是外鍵。
    把一列刪掉，填過那個系的人不會有任何錯誤，他們只是**從分組統計裡消失**，
    而且沒有人會發現。`active=0` 才是正確的做法：留著對得上，只是不進下拉。
    （`OrgDepartment.active` 的原始註解就寫著「停招的留著但不進下拉」。）

ZH: 🔴 **改名一定要連動使用者，否則同一個坑會從另一邊發生。**
    主鍵就是使用者存的那個字串，所以改名 = 把所有填過舊名的人變成孤兒。
    這裡在同一個交易裡一起改，並把「動到幾個人」回報給前端顯示 ——
    管理者按下去之前要看得到影響範圍。

ZH: ⚠ 校區一律對照 `org_seed.CAMPUSES` 驗。打錯的值存進去之後，
    那一列在任何以校區分組的報表裡都會自成一格，看起來像多了一個校區。
    （`crud.set_user_campuses` 也是這樣驗的，兩邊用同一份清單。）
"""
from typing import Any, Optional

import csv
import io
import json
from datetime import datetime, timezone

from fastapi import APIRouter, Body, Depends, File, HTTPException, Query, Response, UploadFile
from fastapi.responses import StreamingResponse
from sqlalchemy import func
from sqlalchemy.orm import Session

from .. import models, org_seed
from ..auth import require_admin
from ..database import get_db

router = APIRouter(tags=["組織對照 Org"])


# ── 共用 ────────────────────────────────────────────────────────────────
def _clean_campus(v: Optional[str]) -> Optional[str]:
    """
    ZH: 校區欄位正規化。空字串一律存成 NULL —— 「沒填」與「填了空字串」
        在報表裡會是兩種不同的分組，看起來像多出一個沒有名字的校區。

    @node job-scheduler/app/routers/org.py::_clean_campus
    """
    c = (v or "").strip()
    if not c:
        return None
    if c not in org_seed.CAMPUSES:
        raise HTTPException(
            status_code=400,
            detail=(f"ZH: 沒有這個校區：{c}（可選：{'、'.join(org_seed.CAMPUSES)}） | "
                    f"EN: Unknown campus: {c}"))
    return c


def _as_active(v: Any) -> int:
    """ZH: 前端可能送 true/1/"1"。一律收斂成 0/1。

    @node job-scheduler/app/routers/org.py::_as_active
    """
    return 0 if v in (0, "0", False, "false", "False") else 1


# ── 學系 ────────────────────────────────────────────────────────────────
@router.get("/departments", summary="學系→學院對照（含每系人數）")
def list_departments(db: Session = Depends(get_db),
                     _: models.User = Depends(require_admin)) -> Any:
    """
    ZH: 含 `users` 人數 —— 停用或改名之前要看得到影響範圍。

    @node job-scheduler/app/routers/org.py::list_departments
    """
    counts = dict(db.query(models.User.department, func.count(models.User.id))
                  .filter(models.User.department.isnot(None))
                  .group_by(models.User.department).all())
    rows = (db.query(models.OrgDepartment)
            .order_by(models.OrgDepartment.college, models.OrgDepartment.name).all())
    return {
        "campuses": org_seed.CAMPUSES,
        "colleges": sorted({r.college for r in rows}),
        "rows": [{"name": r.name, "college": r.college, "campus": r.campus,
                  "name_en": r.name_en or "", "college_en": r.college_en or "",
                  "active": r.active, "users": counts.get(r.name, 0),
                  "source": r.source or "admin"} for r in rows],
    }


@router.put("/departments", summary="批次寫入學系對照")
def save_departments(payload: dict = Body(...),
                     db: Session = Depends(get_db),
                     _: models.User = Depends(require_admin)) -> Any:
    """
    ZH: body: `{"rows": [{"key": 現在的系名 or null(新增), "name", "college",
        "campus", "active"}]}`

    ZH: `key` 與 `name` 不同 = 改名 → 連同 `users.department` 一起改。

    @node job-scheduler/app/routers/org.py::save_departments
    """
    rows = payload.get("rows")
    if not isinstance(rows, list):
        raise HTTPException(status_code=400, detail="ZH: rows 要是陣列 | EN: rows must be a list")

    added = updated = renamed = moved_users = 0
    for r in rows:
        name = (r.get("name") or "").strip()
        college = (r.get("college") or "").strip()
        if not name or not college:
            raise HTTPException(status_code=400,
                                detail="ZH: 系名與學院都不能空白 | EN: name and college are required")
        campus = _clean_campus(r.get("campus"))
        active = _as_active(r.get("active", 1))
        key = (r.get("key") or "").strip() or None
        # ZH: v3.9 英文名。空字串一律存成 None —— 顯示端是用 `or 中文` 退回，
        #     存 "" 與存 None 的行為要一樣，不然「清空」會變成兩種不同的狀態。
        name_en = (r.get("name_en") or "").strip() or None
        college_en = (r.get("college_en") or "").strip() or None

        if key is None:
            if db.get(models.OrgDepartment, name):
                raise HTTPException(status_code=409,
                                    detail=f"ZH: 已經有這個學系：{name} | EN: duplicate: {name}")
            # ZH: v4.9 管理者手動新增的，標 'admin' —— 之後 Alma 重掃時
            #     這一列不會被當成過期的 Alma 資料清掉。
            db.add(models.OrgDepartment(name=name, college=college,
                                        campus=campus, active=active,
                                        name_en=name_en, college_en=college_en,
                                        source='admin'))
            added += 1
            continue

        cur = db.get(models.OrgDepartment, key)
        if not cur:
            raise HTTPException(status_code=404,
                                detail=f"ZH: 找不到學系：{key} | EN: not found: {key}")

        if name != key:
            # ZH: 改名 —— 主鍵換掉，所以是「新增一列 + 刪掉舊列」，
            #     而且**填過舊名的人要一起搬**，否則他們會從分組統計裡消失。
            if db.get(models.OrgDepartment, name):
                raise HTTPException(status_code=409,
                                    detail=f"ZH: 已經有這個學系：{name} | EN: duplicate: {name}")
            moved_users += (db.query(models.User)
                            .filter(models.User.department == key)
                            .update({models.User.department: name},
                                    synchronize_session=False))
            db.add(models.OrgDepartment(name=name, college=college,
                                        campus=campus, active=active,
                                        name_en=name_en, college_en=college_en,
                                        source=cur.source or 'admin'))
            db.delete(cur)
            renamed += 1
            continue

        if ((cur.college, cur.campus, cur.active, cur.name_en, cur.college_en)
                != (college, campus, active, name_en, college_en)):
            cur.college, cur.campus, cur.active = college, campus, active
            cur.name_en, cur.college_en = name_en, college_en
            updated += 1

    db.commit()
    return {"added": added, "updated": updated, "renamed": renamed,
            "moved_users": moved_users}


# ── 行政單位 ────────────────────────────────────────────────────────────
@router.get("/units", summary="行政單位（含每單位人數）")
def list_units(db: Session = Depends(get_db),
               _: models.User = Depends(require_admin)) -> Any:
    """
    ZH: ⚠ 人數用 `users.unit` 比對，而那裡存的是 **path 不是 name**
        （官網底下有兩個「事務組」，名稱會撞 —— 見 models.OrgUnit）。

    @node job-scheduler/app/routers/org.py::list_units
    """
    counts = dict(db.query(models.User.unit, func.count(models.User.id))
                  .filter(models.User.unit.isnot(None))
                  .group_by(models.User.unit).all())
    rows = db.query(models.OrgUnit).order_by(models.OrgUnit.path).all()
    return {
        "campuses": org_seed.CAMPUSES,
        "rows": [{"path": r.path, "name": r.name, "parent": r.parent,
                  "name_en": r.name_en or "",
                  "campus": r.campus, "active": r.active,
                  "users": counts.get(r.path, 0),
                  "source": r.source or "admin"} for r in rows],
    }


@router.put("/units", summary="批次寫入行政單位")
def save_units(payload: dict = Body(...),
               db: Session = Depends(get_db),
               _: models.User = Depends(require_admin)) -> Any:
    """
    ZH: body: `{"rows": [{"key": 現在的 path or null(新增), "name", "parent",
        "campus", "active"}]}`

    ZH: path 由 `parent/name` 推出來，不讓前端直接送 —— 前端自己組的話，
        兩邊的組法遲早會不一致，而不一致的那一天沒有任何錯誤訊息。

    @node job-scheduler/app/routers/org.py::save_units
    """
    rows = payload.get("rows")
    if not isinstance(rows, list):
        raise HTTPException(status_code=400, detail="ZH: rows 要是陣列 | EN: rows must be a list")

    added = updated = renamed = moved_users = 0
    for r in rows:
        name = (r.get("name") or "").strip()
        if not name:
            raise HTTPException(status_code=400,
                                detail="ZH: 單位名稱不能空白 | EN: name is required")
        parent = (r.get("parent") or "").strip() or None
        campus = _clean_campus(r.get("campus"))
        active = _as_active(r.get("active", 1))
        path = f"{parent}/{name}" if parent else name
        key = (r.get("key") or "").strip() or None
        name_en = (r.get("name_en") or "").strip() or None

        if key is None:
            if db.get(models.OrgUnit, path):
                raise HTTPException(status_code=409,
                                    detail=f"ZH: 已經有這個單位：{path} | EN: duplicate: {path}")
            db.add(models.OrgUnit(path=path, name=name, parent=parent,
                                  campus=campus, active=active, name_en=name_en,
                                  source='admin'))
            added += 1
            continue

        cur = db.get(models.OrgUnit, key)
        if not cur:
            raise HTTPException(status_code=404,
                                detail=f"ZH: 找不到單位：{key} | EN: not found: {key}")

        if path != key:
            if db.get(models.OrgUnit, path):
                raise HTTPException(status_code=409,
                                    detail=f"ZH: 已經有這個單位：{path} | EN: duplicate: {path}")
            moved_users += (db.query(models.User)
                            .filter(models.User.unit == key)
                            .update({models.User.unit: path},
                                    synchronize_session=False))
            db.add(models.OrgUnit(path=path, name=name, parent=parent,
                                  campus=campus, active=active, name_en=name_en,
                                  source=cur.source or 'admin'))
            db.delete(cur)
            renamed += 1
            continue

        if (cur.name, cur.campus, cur.active, cur.name_en) != (name, campus, active, name_en):
            cur.name, cur.campus, cur.active = name, campus, active
            cur.name_en = name_en
            updated += 1

    db.commit()
    return {"added": added, "updated": updated, "renamed": renamed,
            "moved_users": moved_users}


# ══════════════════════════════════════════════════════════════════════════
# ZH: 匯出／匯入 —— 讓這張表跟著版控走
# ══════════════════════════════════════════════════════════════════════════
# ZH: 為什麼需要：對照表存在資料庫裡，**不會跟著 repo 走**。
#     換一台機器（或重建 data/）就只剩種子資料，管理者填過的校區、停用、
#     改名全部不見，而且沒有任何提示 —— 只會發現報表的分群變了。
#     匯出成一個 JSON 檔就能進版控，換機器時匯回來。
#
# ZH: 🔴 **匯入只做 upsert，永遠不刪除。**
#     檔案裡沒有的列**原封不動留著**，不會被清掉。
#     理由跟編輯器不給刪同一個：`users.department` 存的是系名本身、沒有外鍵，
#     刪掉一列只會讓填過它的人安靜地從分組統計裡消失。
#     要停用就在檔案裡把那一列的 `active` 設成 0 —— 那是有紀錄、可還原的。
#
# ZH: ⚠ **匯入認不出「改名」。** 檔案以名稱／path 為鍵，所以改過名的列匯進來
#     會被當成**新的一列**，舊的那列還在（而且填過舊名的人仍指著它）。
#     這是知情的取捨：要認得改名就得在檔案裡記 id，而這兩張表刻意用名稱當主鍵
#     （見 models.OrgDepartment）。改名請用管理端的編輯器，那裡會連動使用者。
#     匯入的預覽會把「新增」列出來，看到不該新增的名字就知道是這個情況。

ORG_EXPORT_VERSION = 2   # ZH: v3.9 多了 name_en / college_en


# ZH: 匯出的兩種用途，不要混在一起：
#       · **JSON** ＝ 給機器的：它同時是**匯入格式**（/import 的 body 就是這份檔），
#         所以欄位名與值都照資料庫原樣（active 是 0/1）。換機器要救回這張表靠它。
#       · **CSV / XLSX** ＝ 給人看的：欄位名是中文、`active` 寫成「是／否」、
#         多一欄「來源」讓人分得出哪幾列是 Alma 來的、哪幾列是人工加的。
#         這兩種**也匯得回來**（`/import-file`），但有一個差別要講明：
#         🔴 試算表的**空格一律當成「別動」**，所以匯入永遠不會清空欄位。
#
# ZH: CSV 沒有「工作表」的概念，所以學系與行政單位**併成一張表**，
#     第一欄「類型」分辨。XLSX 則拆成兩個工作表（那是它天生該做的事）。
_ORG_CSV_HEADERS = ["類型", "名稱", "英文名稱", "上層", "上層英文",
                    "完整路徑", "校區", "啟用", "來源"]
_ORG_DEPT_HEADERS = ["學系名稱", "英文名稱", "學院", "學院英文", "校區", "啟用", "來源"]
_ORG_UNIT_HEADERS = ["完整路徑", "名稱", "英文名稱", "上層", "校區", "啟用", "來源"]


def _yn(v) -> str:
    """ZH: 給人看的啟用欄。0/1 在 Excel 裡看起來像代碼，而這一欄是要用眼睛掃的。

    @node job-scheduler/app/routers/org.py::_yn
    """
    return "是" if v else "否"


def _org_rows(depts, units):
    """ZH: CSV 用的單一張表（學系與行政單位併在一起，第一欄分辨）。

    ZH: 行政單位沒有「上層英文」可填 —— 資料庫只存每一列自己的 `name_en`，
        上層是誰是用 path 串出來的。這裡留空而不是拼一個出來：
        拼出來的英文名沒有人維護過，看起來卻像是正式的。

    @node job-scheduler/app/routers/org.py::_org_rows
    """
    rows = []
    for d in depts:
        rows.append(["學系", d.name, d.name_en or "", d.college or "", d.college_en or "",
                     "", d.campus or "", _yn(d.active), d.source or ""])
    for u in units:
        rows.append(["行政單位", u.name, u.name_en or "", u.parent or "", "",
                     u.path, u.campus or "", _yn(u.active), u.source or ""])
    return rows


def _org_csv(depts, units, stamp: str):
    """ZH: CSV。**BOM 不能省** —— 沒有它 Excel 會用系統編碼開，中文全變亂碼
            （與 admin.export_users 同一個理由，那邊也加了）。

    @node job-scheduler/app/routers/org.py::_org_csv
    """
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(_ORG_CSV_HEADERS)
    w.writerows(_org_rows(depts, units))
    content = ("﻿" + buf.getvalue()).encode("utf-8")
    return StreamingResponse(
        io.BytesIO(content),
        # ZH: 只寫 "text/csv" —— Starlette 對 text/* 會自己補 charset，
        #     自己再寫一次會變成 `text/csv; charset=utf-8; charset=utf-8`。
        media_type="text/csv",
        headers={"Content-Disposition": f'attachment; filename="org-mapping-{stamp}.csv"'},
    )


def _org_xlsx(depts, units, stamp: str):
    """ZH: XLSX。學系與行政單位各一個工作表 —— 兩者的欄位本來就不一樣，
            併成一張表會有一半的格子是空的。

    @node job-scheduler/app/routers/org.py::_org_xlsx
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise HTTPException(status_code=500,
                            detail="ZH: openpyxl 未安裝 | EN: openpyxl not installed")

    wb = Workbook()
    sheets = [
        ("學系", _ORG_DEPT_HEADERS,
         [[d.name, d.name_en or "", d.college or "", d.college_en or "",
           d.campus or "", _yn(d.active), d.source or ""] for d in depts]),
        ("行政單位", _ORG_UNIT_HEADERS,
         [[u.path, u.name, u.name_en or "", u.parent or "",
           u.campus or "", _yn(u.active), u.source or ""] for u in units]),
    ]
    for i, (title, headers, rows) in enumerate(sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        ws.title = title
        ws.append(headers)
        for r in rows:
            ws.append(r)
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="DDEBF7")
        ws.freeze_panes = "A2"
        for col_idx, col_name in enumerate(headers, start=1):
            longest = max([len(str(col_name))]
                          + [len(str(r[col_idx - 1])) for r in rows[:200]] or [0])
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width =                 min(longest + 2, 40)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="org-mapping-{stamp}.xlsx"'},
    )


@router.get("/export", summary="匯出組織對照表（json 可匯回／csv・xlsx 給人看）")
def export_org(fmt: str = Query("json", pattern="^(json|csv|xlsx)$",
                                description="ZH: json（可匯回）／csv／xlsx"),
               db: Session = Depends(get_db),
               _: models.User = Depends(require_admin)):
    """
    ZH: 回一個可下載的檔案。**不含人數** —— 那是衍生資料，
        帶著它會讓兩台機器的檔案內容不同而看起來像有差異。

    ZH: `fmt=json`（預設，維持舊行為）＝ 版控用（原始值、可 diff）；
        `fmt=csv`、`fmt=xlsx` ＝ 給人用 Excel 看的。三種都能經 `/import-file` 匯回，
        差別是試算表的空格＝「別動」（見那一段的說明）。

    @node job-scheduler/app/routers/org.py::export_org
    """
    depts = (db.query(models.OrgDepartment)
             .order_by(models.OrgDepartment.college, models.OrgDepartment.name).all())
    units = db.query(models.OrgUnit).order_by(models.OrgUnit.path).all()

    if fmt in ("csv", "xlsx"):
        stamp = datetime.now(timezone.utc).strftime("%Y%m%d_%H%M%S")
        return (_org_csv(depts, units, stamp) if fmt == "csv"
                else _org_xlsx(depts, units, stamp))

    body = {
        "_說明": [
            "組織對照表（學系→學院、行政單位）。這是給版控用的匯出檔。",
            "匯回的方式：管理端 → 平台設定 → 組織對照表 → 匯入。",
            "⚠ 匯入只會新增與更新，**不會刪除**檔案裡沒有的列。",
            "⚠ 匯入認不出改名 —— 改過名的會變成新的一列。改名請用編輯器。",
        ],
        "version": ORG_EXPORT_VERSION,
        "exported_at": datetime.now(timezone.utc).isoformat(),
        "campuses": org_seed.CAMPUSES,
        "departments": [{"name": d.name, "college": d.college,
                         "campus": d.campus, "active": d.active,
                         "name_en": d.name_en or "",
                         "college_en": d.college_en or ""} for d in depts],
        "units": [{"path": u.path, "name": u.name, "parent": u.parent,
                   "campus": u.campus, "active": u.active,
                   "name_en": u.name_en or ""} for u in units],
    }
    return Response(
        content=json.dumps(body, ensure_ascii=False, indent=2),
        media_type="application/json; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="org-mapping.json"'},
    )


def _apply_org_import(db: Session, payload: dict, dry_run: bool) -> Any:
    """
    ZH: 匯入的本體。JSON 與試算表兩個入口都走這裡 ——
        解析格式是兩回事，**套用規則只能有一份**（兩份遲早一邊有防呆一邊沒有）。

    ZH: 🔴 預設是預覽而不是直接寫 —— 這張表牽動全站的分群統計，
        「按錯就套用」與「按錯先給你看」的代價差很多。

    @node job-scheduler/app/routers/org.py::_apply_org_import
    """
    # ZH: v3.9 —— 版本改成「不高於現在」就收，不再要求相等。
    #     格式是**往上相容的疊加**（v2 只是多了 name_en / college_en），
    #     舊檔缺鍵時下面會保留現有的英文名而不是清掉它。
    #     嚴格相等的話，bump 一次版本就會讓所有進版控的舊匯出檔失效 ——
    #     而那些檔案的內容其實完全還能用。
    #     比現在**新**的仍然要擋：那是未來的格式，我們不知道它多了什麼。
    ver = payload.get("version")
    if ver is not None and (not isinstance(ver, int) or ver > ORG_EXPORT_VERSION):
        raise HTTPException(
            status_code=400,
            detail=(f"ZH: 檔案版本 {ver} 比這個平台認得的 {ORG_EXPORT_VERSION} 還新 | "
                    f"EN: unsupported export version {ver}"))

    depts = payload.get("departments")
    units = payload.get("units")
    if not isinstance(depts, list) or not isinstance(units, list):
        raise HTTPException(
            status_code=400,
            detail="ZH: 檔案裡要有 departments 與 units 兩個陣列 | EN: need both arrays")

    report = {"dry_run": dry_run,
              "departments": {"added": [], "updated": [], "unchanged": 0},
              "units": {"added": [], "updated": [], "unchanged": 0},
              "untouched_in_db": {"departments": 0, "units": 0}}

    seen_d, seen_u = set(), set()

    for r in depts:
        name = (r.get("name") or "").strip()
        college = (r.get("college") or "").strip()
        if not name or not college:
            raise HTTPException(status_code=400,
                                detail=f"ZH: 學系缺欄位：{r} | EN: bad department row")
        # ZH: v4.12 `campus` / `active` 也照 name_en 的規矩：**缺鍵＝保留現值**。
        #     這是為了試算表匯入 —— 試算表沒有「省略欄位」這回事，欄一定在、
        #     只有格子是空的，所以解析器把空格**不放進 dict**，語意才接得上。
        #     若不這樣：空白校區會清掉現有校區，空白啟用會把停用的系**重新打開**，
        #     兩者都是安靜地發生。新增的列仍用預設（校區空、啟用）。
        has_campus = "campus" in r
        has_active = "active" in r
        campus = _clean_campus(r.get("campus"))
        active = _as_active(r.get("active", 1))
        # ZH: v3.9 英文名。舊版匯出檔沒有這兩個鍵 → 會是 None，
        #     於是匯入舊檔會把既有的英文名**清掉**。所以缺鍵時保留現值，
        #     只有檔案裡明確給了才覆寫（給空字串＝明確要清空）。
        has_en = "name_en" in r
        has_cen = "college_en" in r
        name_en = ((r.get("name_en") or "").strip() or None) if has_en else None
        college_en = ((r.get("college_en") or "").strip() or None) if has_cen else None
        seen_d.add(name)
        cur = db.get(models.OrgDepartment, name)
        if cur is None:
            report["departments"]["added"].append(name)
            if not dry_run:
                db.add(models.OrgDepartment(name=name, college=college,
                                            campus=campus, active=active,
                                            name_en=name_en, college_en=college_en,
                                            source='admin'))
        else:
            want_en = name_en if has_en else cur.name_en
            want_cen = college_en if has_cen else cur.college_en
            campus = campus if has_campus else cur.campus
            active = active if has_active else cur.active
            if ((cur.college, cur.campus, cur.active, cur.name_en, cur.college_en)
                    != (college, campus, active, want_en, want_cen)):
                report["departments"]["updated"].append(name)
                if not dry_run:
                    cur.college, cur.campus, cur.active = college, campus, active
                    cur.name_en, cur.college_en = want_en, want_cen
            else:
                report["departments"]["unchanged"] += 1

    for r in units:
        name = (r.get("name") or "").strip()
        if not name:
            raise HTTPException(status_code=400,
                                detail=f"ZH: 單位缺名稱：{r} | EN: bad unit row")
        parent = (r.get("parent") or "").strip() or None
        # ZH: path 一律重算，不採用檔案裡的 —— 檔案可能是別的版本組出來的。
        path = f"{parent}/{name}" if parent else name
        has_campus = "campus" in r       # ZH: 理由同學系那段（試算表的空格＝別動）
        has_active = "active" in r
        campus = _clean_campus(r.get("campus"))
        active = _as_active(r.get("active", 1))
        has_en = "name_en" in r          # ZH: 理由同上：舊檔缺鍵時不要清掉現值
        name_en = ((r.get("name_en") or "").strip() or None) if has_en else None
        seen_u.add(path)
        cur = db.get(models.OrgUnit, path)
        if cur is None:
            report["units"]["added"].append(path)
            if not dry_run:
                db.add(models.OrgUnit(path=path, name=name, parent=parent,
                                      campus=campus, active=active, name_en=name_en,
                                      source='admin'))
        else:
            want_en = name_en if has_en else cur.name_en
            campus = campus if has_campus else cur.campus
            active = active if has_active else cur.active
            if (cur.name, cur.campus, cur.active, cur.name_en) != (name, campus, active, want_en):
                report["units"]["updated"].append(path)
                if not dry_run:
                    cur.name, cur.campus, cur.active = name, campus, active
                    cur.name_en = want_en
            else:
                report["units"]["unchanged"] += 1

    # ZH: 資料庫有、檔案沒有的 —— 只回報數量，**不動它們**。
    #     這個數字是給人看的訊號：不是 0 就表示兩邊不同步，可能是檔案舊了，
    #     也可能是有人改過名（改名會在上面變成「新增」）。
    report["untouched_in_db"]["departments"] = sum(
        1 for (n,) in db.query(models.OrgDepartment.name).all() if n not in seen_d)
    report["untouched_in_db"]["units"] = sum(
        1 for (p,) in db.query(models.OrgUnit.path).all() if p not in seen_u)

    if dry_run:
        db.rollback()
    else:
        db.commit()
    return report


# ══════════════════════════════════════════════════════════════════════════
# ZH: 匯入的三種入口。**套用規則只有 `_apply_org_import` 一份**，
#     這一層只負責「把檔案變成 payload」。
#
# ZH: 🔴 試算表沒有「省略欄位」這回事 —— 欄一定在，只有格子是空的。
#     所以解析器遇到空格就**不把那個鍵放進 dict**，讓它落在
#     `_apply_org_import` 的「缺鍵＝保留現值」那條路上。
#     這是整段最重要的一個決定：不這樣做的話，一個沒填英文名的 Excel
#     匯進來會把全校的英文名清空，而且畫面上不會有任何錯誤。
#     代價講明白：**匯入無法用來清空欄位**，要清空請用管理端的編輯器。
# ══════════════════════════════════════════════════════════════════════════
_ORG_TRUE = {"1", "是", "y", "yes", "true", "啟用", "on"}
_ORG_FALSE = {"0", "否", "n", "no", "false", "停用", "off"}


def _cell(v) -> str:
    """ZH: 儲存格 → 去頭尾空白的字串。None／空白都變空字串。

    ZH: openpyxl 會把「1」這種格子給成數值，而 str(1.0) 是 "1.0" ——
        那會讓「啟用」與代碼類欄位長出一個小數點。

    @node job-scheduler/app/routers/org.py::_cell
    """
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def _active_cell(v):
    """ZH: 啟用欄 → 1／0／None（None＝空白＝別動）。認不得的值要吵，不要猜。

    @node job-scheduler/app/routers/org.py::_active_cell
    """
    t = _cell(v)
    if not t:
        return None
    low = t.lower()
    if low in _ORG_TRUE:
        return 1
    if low in _ORG_FALSE:
        return 0
    raise HTTPException(
        status_code=400,
        detail="ZH: 啟用欄看不懂「%s」，請填「是」或「否」 | EN: bad active value" % t)


def _put(d: dict, key: str, value: str):
    """ZH: 只有非空才放進去 —— 空的就讓它「缺鍵」。見本段開頭的說明。

    @node job-scheduler/app/routers/org.py::_put
    """
    if value:
        d[key] = value


def _row_to_payload(kind: str, get, depts: list, units: list):
    """ZH: 一列 → 學系或行政單位。`get` 是「欄名 → 值」的取值函式。

    @node job-scheduler/app/routers/org.py::_row_to_payload
    """
    name = _cell(get("名稱"))
    if not name:
        return                          # ZH: 整列空白（試算表尾巴常有）→ 跳過
    active = _active_cell(get("啟用"))
    if kind == "學系":
        row = {"name": name, "college": _cell(get("上層"))}
        _put(row, "name_en", _cell(get("英文名稱")))
        _put(row, "college_en", _cell(get("上層英文")))
        _put(row, "campus", _cell(get("校區")))
        if active is not None:
            row["active"] = active
        depts.append(row)
    else:
        row = {"name": name}
        _put(row, "name_en", _cell(get("英文名稱")))
        _put(row, "parent", _cell(get("上層")))
        _put(row, "campus", _cell(get("校區")))
        if active is not None:
            row["active"] = active
        units.append(row)


def _norm_headers(headers: list) -> list:
    """ZH: 表頭正規化 —— xlsx 兩張工作表用的是各自的欄名（「學系名稱」「學院」…），
        CSV 併成一張用的是通用欄名（「名稱」「上層」…）。統一成後者再解析，
        才不會有兩份幾乎一樣的欄位對應表。

    @node job-scheduler/app/routers/org.py::_norm_headers
    """
    out = []
    for h in headers:
        h = _cell(h)
        if h in ("學系名稱", "單位名稱"):
            h = "名稱"
        elif h in ("學院", "上層單位"):
            h = "上層"
        elif h == "學院英文":
            h = "上層英文"
        out.append(h)
    return out


def _payload_from_csv(raw: bytes) -> dict:
    """ZH: CSV → payload。第一欄「類型」分辨學系／行政單位（匯出時就是這樣併的）。

    @node job-scheduler/app/routers/org.py::_payload_from_csv
    """
    # ZH: utf-8-sig 才吃得下自己匯出的 BOM；再退一步試 cp950
    #     （有人會在 Excel 裡「另存新檔 → CSV」，那在中文 Windows 上是 cp950）。
    text = None
    for enc in ("utf-8-sig", "cp950"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise HTTPException(status_code=400,
                            detail="ZH: CSV 編碼看不懂，請存成 UTF-8 | EN: undecodable CSV")
    rows = list(csv.reader(io.StringIO(text)))
    if not rows:
        raise HTTPException(status_code=400, detail="ZH: 檔案是空的 | EN: empty file")
    headers = _norm_headers(rows[0])
    if "類型" not in headers or "名稱" not in headers:
        raise HTTPException(
            status_code=400,
            detail="ZH: CSV 要有「類型」與「名稱」欄（請用本頁匯出的 CSV 當範本） | "
                   "EN: CSV needs 類型 and 名稱 columns")
    depts = []
    units = []
    for r in rows[1:]:
        def get(col, _r=r):
            """@node job-scheduler/app/routers/org.py::_payload_from_csv.<get>"""
            if col not in headers:
                return ""
            i = headers.index(col)
            return _r[i] if i < len(_r) else ""
        kind = _cell(get("類型"))
        if not kind and not _cell(get("名稱")):
            continue                     # ZH: 空白列
        if kind not in ("學系", "行政單位"):
            raise HTTPException(
                status_code=400,
                detail="ZH: 類型只能是「學系」或「行政單位」，看到「%s」 | "
                       "EN: bad 類型" % kind)
        _row_to_payload(kind, get, depts, units)
    return {"departments": depts, "units": units}


def _payload_from_xlsx(raw: bytes) -> dict:
    """ZH: XLSX → payload。工作表名就是類型（匯出時分成「學系」「行政單位」兩張）。

    @node job-scheduler/app/routers/org.py::_payload_from_xlsx
    """
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise HTTPException(status_code=500,
                            detail="ZH: openpyxl 未安裝 | EN: openpyxl not installed")
    try:
        wb = load_workbook(io.BytesIO(raw), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400,
                            detail="ZH: 這個檔案不是有效的 Excel（.xlsx）：%s | EN: bad xlsx" % e)
    known = [n for n in wb.sheetnames if n in ("學系", "行政單位")]
    if not known:
        raise HTTPException(
            status_code=400,
            detail="ZH: 找不到「學系」或「行政單位」工作表（請用本頁匯出的 Excel 當範本） | "
                   "EN: need 學系 / 行政單位 sheets")
    depts = []
    units = []
    for kind in known:
        ws = wb[kind]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        headers = _norm_headers(list(rows[0]))
        if "名稱" not in headers:
            raise HTTPException(
                status_code=400,
                detail="ZH: 工作表「%s」找不到名稱欄 | EN: no name column" % kind)
        for r in rows[1:]:
            def get(col, _r=r, _h=headers):
                """@node job-scheduler/app/routers/org.py::_payload_from_xlsx.<get>"""
                if col not in _h:
                    return ""
                i = _h.index(col)
                return _r[i] if i < len(_r) else ""
            _row_to_payload(kind, get, depts, units)
    return {"departments": depts, "units": units}


@router.post("/import", summary="匯入組織對照表（JSON；預設先預覽）")
def import_org(payload: dict = Body(...),
               dry_run: bool = True,
               db: Session = Depends(get_db),
               _: models.User = Depends(require_admin)) -> Any:
    """
    ZH: body 就是 JSON 匯出檔的內容。`dry_run=true`（預設）只回報會發生什麼，不寫入。

    @node job-scheduler/app/routers/org.py::import_org
    """
    return _apply_org_import(db, payload, dry_run)


@router.post("/import-file", summary="匯入組織對照表（.xlsx／.csv／.json；預設先預覽）")
async def import_org_file(file: UploadFile = File(...),
                          dry_run: bool = True,
                          db: Session = Depends(get_db),
                          _: models.User = Depends(require_admin)) -> Any:
    """
    ZH: 上傳匯出檔（副檔名決定怎麼解析），預設一樣只預覽。

    ZH: 🔴 **匯入不會清空欄位**：試算表的空格一律當成「別動」
        （理由見本段開頭那塊註解）。要清空請用管理端的編輯器。

    @node job-scheduler/app/routers/org.py::import_org_file
    """
    name = (file.filename or "").lower()
    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=400, detail="ZH: 檔案是空的 | EN: empty file")
    if name.endswith(".xlsx"):
        payload = _payload_from_xlsx(raw)
    elif name.endswith(".csv"):
        payload = _payload_from_csv(raw)
    elif name.endswith(".json"):
        try:
            payload = json.loads(raw.decode("utf-8-sig"))
        except Exception as e:
            raise HTTPException(status_code=400,
                                detail="ZH: JSON 讀不出來：%s | EN: bad json" % e)
    else:
        raise HTTPException(
            status_code=400,
            detail="ZH: 只收 .xlsx / .csv / .json | EN: only .xlsx / .csv / .json")
    return _apply_org_import(db, payload, dry_run)
